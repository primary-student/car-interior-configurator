"""把生成的 pptx 反向渲染成 PNG，用于版式自检（无 Office / LibreOffice 环境下的替代方案）。

按 python-pptx 读出的真实坐标、字号、颜色绘制；中文换行用 PIL 的 font.getlength 逐字测量。
"""
import os
import sys
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

HERE = os.path.dirname(os.path.abspath(__file__))
PPTX = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "内秀NEIXIU_AWS鬼点子计划_参赛PPT.pptx")
OUT = os.path.join(HERE, "preview")
os.makedirs(OUT, exist_ok=True)

EMU_IN = 914400.0
W_PX, H_PX = 1920, 1080
W_IN = 13.333                      # 幻灯片宽（英寸）
PX_PER_IN = W_PX / W_IN            # 144 px/in

FONT_FILES = {
    "Microsoft YaHei Light": ["C:/Windows/Fonts/msyhl.ttc", "C:/Windows/Fonts/msyh.ttc"],
    "Microsoft YaHei": ["C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/msyhbd.ttc"],
}
_cache = {}


def font(name, px, bold=False):
    key = (name, px, bold)
    if key in _cache:
        return _cache[key]
    for path in FONT_FILES.get(name, ["C:/Windows/Fonts/msyh.ttc"]):
        if os.path.exists(path):
            try:
                f = ImageFont.truetype(path, px, index=0)
                _cache[key] = f
                return f
            except Exception:
                continue
    f = ImageFont.load_default()
    _cache[key] = f
    return f


def px(v):
    return v / EMU_IN * PX_PER_IN


def rgb(color_format, default=(255, 255, 255)):
    try:
        return tuple(color_format.rgb)
    except Exception:
        return default


def draw_text(d, box_xy, w, h, shape):
    tf = shape.text_frame
    x, y = box_xy
    cy = y
    for p in tf.paragraphs:
        if not p.runs:
            continue
        runs = p.runs
        size_pt = runs[0].font.size.pt if runs[0].font.size else 15
        bold = bool(runs[0].font.bold)
        fname = runs[0].font.name or "Microsoft YaHei Light"
        color = rgb(runs[0].font.color, (242, 236, 229))
        size_px = int(round(size_pt * 2.0))          # 144dpi: 1pt = 2px
        f = font(fname, size_px, bold)
        spacing = p.line_spacing if isinstance(p.line_spacing, float) else 1.2
        # 逐字换行
        lines, cur = [], ""
        for ch in runs[0].text:
            trial = cur + ch
            if f.getlength(trial) > w + 0.5 and cur:
                lines.append(cur)
                cur = ch
            else:
                cur = trial
        lines.append(cur)
        line_h = size_px * spacing
        for ln in lines:
            lw = f.getlength(ln)
            if p.alignment == PP_ALIGN.CENTER:
                tx = x + (w - lw) / 2
            elif p.alignment == PP_ALIGN.RIGHT:
                tx = x + w - lw
            else:
                tx = x
            d.text((tx, cy), ln, font=f, fill=color)
            cy += line_h


def render(slide, idx):
    img = Image.new("RGB", (W_PX, H_PX), (13, 13, 15))
    d = ImageDraw.Draw(img)
    for sh in slide.shapes:
        x, y = px(sh.left or 0), px(sh.top or 0)
        w, h = px(sh.width or 0), px(sh.height or 0)
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            from io import BytesIO
            fname = getattr(sh.image, "filename", None)
            src = os.path.join(HERE, "assets", os.path.basename(fname)) if fname else ""
            if src and os.path.isfile(src):
                im = Image.open(src).convert("RGB")
            else:
                im = Image.open(BytesIO(sh.image.blob)).convert("RGB")
            im = im.resize((max(1, int(w)), max(1, int(h))), Image.LANCZOS)
            img.paste(im, (int(x), int(y)))
            continue
        if sh.has_text_frame and sh.text_frame.text.strip():
            draw_text(d, (x, y), w, h, sh)
            continue
        # autoshape
        try:
            fill = rgb(sh.fill.fore_color)
        except Exception:
            fill = None
        line = None
        try:
            if sh.line.width:
                line = rgb(sh.line.color)
        except Exception:
            line = None
        d.rectangle([x, y, x + w, y + h], fill=fill, outline=line, width=1)
    path = os.path.join(OUT, f"s{idx:02d}.png")
    img.save(path)
    return path


prs = Presentation(PPTX)
paths = [render(s, i + 1) for i, s in enumerate(prs.slides)]
print(f"{len(paths)} 页 ->", os.path.join(OUT))
